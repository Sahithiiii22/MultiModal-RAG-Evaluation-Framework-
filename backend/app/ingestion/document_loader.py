import os
import uuid
import re
import pandas as pd
from typing import List, Dict, Any
from app.models.schemas import DocumentMetadata, DocumentChunk
from app.services.multimodal_service import multimodal_service

class DocumentLoader:
    """Parses Multimodal documents: PDF, TXT, DOCX, CSV, Image, Video, and Audio into standard document structures."""
    
    @staticmethod
    def load_document(file_path: str, filename: str) -> Dict[str, Any]:
        ext = os.path.splitext(filename)[1].lower()
        doc_id = str(uuid.uuid4())
        file_size = os.path.getsize(file_path) if os.path.exists(file_path) else 0
        
        pages = []
        file_type = ext.replace(".", "")
        
        if ext == ".pdf":
            pages = DocumentLoader._load_pdf(file_path, filename)
        elif ext in [".txt", ".md", ".log", ".json", ".rst"]:
            pages = DocumentLoader._load_txt(file_path)
        elif ext in [".docx", ".doc"]:
            pages = DocumentLoader._load_docx(file_path, filename)
        elif ext in [".csv", ".tsv"]:
            pages = DocumentLoader._load_csv(file_path)
        elif ext in [".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff", ".gif", ".ico", ".svg"]:
            pages = multimodal_service.process_image(file_path, filename)
        elif ext in [".mp4", ".mov", ".avi", ".mkv", ".webm", ".flv", ".wmv", ".m4v"]:
            pages = multimodal_service.process_video(file_path, filename)
        elif ext in [".mp3", ".wav", ".m4a", ".ogg", ".flac", ".aac", ".wma"]:
            pages = multimodal_service.process_audio(file_path, filename)
        else:
            pages = DocumentLoader._load_txt(file_path)
            
        # Ensure at least 1 page exists
        if not pages:
            pages = [{
                "page_number": 1,
                "text": f"Document: {filename} (Type: {file_type.upper()}). File size: {file_size} bytes.",
                "source_type": file_type,
                "content_type": "text"
            }]
            
        metadata = DocumentMetadata(
            doc_id=doc_id,
            filename=filename,
            file_type=file_type,
            file_size_bytes=file_size,
            upload_timestamp=pd.Timestamp.now().isoformat(),
            chunk_count=0,
            page_count=len(pages)
        )
        
        return {
            "metadata": metadata,
            "pages": pages
        }

    @staticmethod
    def _load_pdf(file_path: str, filename: str) -> List[Dict[str, Any]]:
        pages = []
        try:
            import pypdf
            reader = pypdf.PdfReader(file_path)
            for idx, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                # Clean up null bytes and strange artifacts
                text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', text).strip()
                if not text:
                    # Fallback description for scanned / image-based page
                    text = f"[Scanned / Graphic PDF Page] File: {filename} (Page {idx + 1}). Visual image page with diagrammatic or identity layout."
                
                pages.append({
                    "page_number": idx + 1,
                    "text": text,
                    "source_type": "pdf",
                    "content_type": "text"
                })
        except Exception as e:
            try:
                with open(file_path, "rb") as f:
                    content = f.read().decode("utf-8", errors="ignore")
                    clean_content = re.sub(r'[^\x20-\x7E\n\r\t]', ' ', content).strip()
                    if clean_content and len(clean_content) > 30:
                        pages.append({
                            "page_number": 1,
                            "text": clean_content[:3000],
                            "source_type": "pdf",
                            "content_type": "text"
                        })
                    else:
                        pages.append({
                            "page_number": 1,
                            "text": f"[PDF Document] File: {filename}. Uploaded PDF file with graphic content.",
                            "source_type": "pdf",
                            "content_type": "text"
                        })
            except Exception:
                pages.append({
                    "page_number": 1,
                    "text": f"[PDF Document] File: {filename}. Uploaded document.",
                    "source_type": "pdf",
                    "content_type": "text"
                })
                
        return pages if pages else [{
            "page_number": 1,
            "text": f"[PDF Document] File: {filename}.",
            "source_type": "pdf",
            "content_type": "text"
        }]

    @staticmethod
    def _load_txt(file_path: str) -> List[Dict[str, Any]]:
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
        except Exception:
            content = ""
        return [{
            "page_number": 1,
            "text": content.strip() or "[Empty Text File]",
            "source_type": "txt",
            "content_type": "text"
        }]

    @staticmethod
    def _load_docx(file_path: str, filename: str) -> List[Dict[str, Any]]:
        try:
            import docx
            doc = docx.Document(file_path)
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            tables_text = []
            for table in doc.tables:
                for row in table.rows:
                    row_vals = [c.text.strip() for c in row.cells if c.text.strip()]
                    if row_vals:
                        tables_text.append(" | ".join(row_vals))
            
            combined = "\n".join(paragraphs + tables_text)
            return [{
                "page_number": 1,
                "text": combined.strip() or f"[DOCX File: {filename}]",
                "source_type": "docx",
                "content_type": "text"
            }]
        except Exception:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            return [{
                "page_number": 1,
                "text": content.strip() or f"[DOCX File: {filename}]",
                "source_type": "docx",
                "content_type": "text"
            }]

    @staticmethod
    def _load_csv(file_path: str) -> List[Dict[str, Any]]:
        try:
            df = pd.read_csv(file_path)
            summary = f"CSV Document with columns: {', '.join(df.columns.tolist())}. Total rows: {len(df)}.\n\n"
            rows_text = []
            for i, row in df.iterrows():
                row_str = " | ".join([f"{col}: {val}" for col, val in row.items()])
                rows_text.append(f"Row {i+1}: {row_str}")
            
            pages = []
            chunk_size = 50
            for page_idx in range(0, len(rows_text), chunk_size):
                page_rows = rows_text[page_idx : page_idx + chunk_size]
                pages.append({
                    "page_number": (page_idx // chunk_size) + 1,
                    "text": summary + "\n".join(page_rows),
                    "source_type": "csv",
                    "content_type": "table"
                })
            return pages if pages else [{"page_number": 1, "text": summary, "source_type": "csv", "content_type": "table"}]
        except Exception as e:
            return [{"page_number": 1, "text": f"CSV File load error: {str(e)}", "source_type": "csv", "content_type": "text"}]
